import io

import pytest
from fastapi.testclient import TestClient

from app.main import SESSIONS, app
from reconiq import __version__
from reconiq.releases import RELEASES


@pytest.fixture()
def client():
    with TestClient(app) as test_client:
        yield test_client


@pytest.fixture()
def session(client):
    response = client.post("/api/session")
    assert response.status_code == 200
    return {"X-Session": response.json()["session_id"]}


def test_index_and_static_assets(client):
    page = client.get("/")
    assert page.status_code == 200
    assert 'id="tab-whats-new"' in page.text
    assert 'aria-controls="view-whats-new"' in page.text
    assert 'id="view-whats-new"' in page.text
    assert client.get("/static/styles.css").status_code == 200
    assert client.get("/static/app.js").status_code == 200


@pytest.mark.parametrize("headers", [{}, {"X-Session": "not-a-session"}])
def test_release_history_is_public_without_creating_sessions(client, headers):
    sessions_before = set(SESSIONS)
    response = client.get("/api/releases", headers=headers)
    assert response.status_code == 200
    assert response.json() == {"current_version": __version__, "releases": RELEASES}
    assert set(SESSIONS) == sessions_before


def test_quality_tab_is_in_app_shell(client):
    page = client.get("/")
    assert page.status_code == 200
    assert 'id="tab-quality"' in page.text
    assert 'id="view-quality"' in page.text
    assert 'id="quality-toggles"' in page.text


def test_quality_catalogue_is_public_and_describes_test_evidence(client):
    sessions_before = set(SESSIONS)
    response = client.get("/api/quality")
    assert response.status_code == 200
    quality = response.json()
    assert quality["verification"]["passed"] == 108
    assert quality["verification"]["failed"] == 0
    assert quality["verification"]["command"] == "python -m pytest tests -q"
    assert {suite["id"] for suite in quality["suites"]} >= {
        "normalization", "schema", "ingestion", "reconciliation", "api", "release-history"
    }
    reconciliation = next(s for s in quality["suites"] if s["id"] == "reconciliation")
    assert "test_demo_dataset_surfaces_the_planted_exceptions" in reconciliation["scripts"]
    assert reconciliation["steps"] and reconciliation["impact"]
    assert quality["incidents"] == []
    assert set(SESSIONS) == sessions_before


def test_quality_catalogue_is_read_only(client):
    assert client.post("/api/quality").status_code == 405


def test_release_history_does_not_change_reconciliation_work(client, session):
    client.post("/api/demo/load", headers=session)
    result = client.post("/api/reconcile", headers=session, json={})
    assert result.status_code == 200
    stored = SESSIONS[session["X-Session"]]
    sources_before = dict(stored.sources)
    result_before = stored.result

    assert client.get("/api/releases", headers=session).status_code == 200
    assert stored.sources == sources_before
    assert stored.result is result_before


def test_application_version_is_consistent_in_api_metadata(client):
    assert client.get("/api/health").json()["version"] == __version__
    assert client.get("/openapi.json").json()["info"]["version"] == __version__


@pytest.mark.parametrize("method", ["post", "put", "delete"])
def test_release_history_is_read_only(client, method):
    assert getattr(client, method)("/api/releases").status_code == 405


def test_demo_end_to_end(client, session):
    demo = client.get("/api/demo").json()
    assert len(demo["files"]) == 4

    loaded = client.post("/api/demo/load", headers=session)
    assert loaded.status_code == 200
    sources = loaded.json()["sources"]
    assert len(sources) == 4
    assert all(source["ready"] for source in sources)

    result = client.post("/api/reconcile", headers=session, json={"mappings": [], "config": {}})
    assert result.status_code == 200
    payload = result.json()
    assert payload["kpis"]["exceptions"] == 18
    assert payload["kpis"]["matched"] == 34
    assert len(payload["findings"]) == 18
    assert payload["findings"][0]["severity"] == "high"
    assert payload["assumptions"] and payload["data_quality"]
    # statement findings inherit the insurer attributed to the statement
    assert all(f["insurer"] for f in payload["findings"] if f["source_id"] != "broker")


def test_exports_are_valid_files(client, session):
    client.post("/api/demo/load", headers=session)
    client.post("/api/reconcile", headers=session, json={"mappings": [], "config": {}})

    csv_export = client.get("/api/export/csv", headers=session)
    assert csv_export.status_code == 200
    assert csv_export.headers["content-disposition"].endswith('.csv"')
    assert b"Duplicate on statement" in csv_export.content

    xlsx_export = client.get("/api/export/xlsx", headers=session)
    assert xlsx_export.status_code == 200
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(xlsx_export.content))
    assert workbook.sheetnames == ["Exceptions", "Summary", "Matched pairs"]
    assert workbook["Exceptions"].max_row == 19  # 18 findings + header


def test_export_before_reconciling_is_rejected(client, session):
    response = client.get("/api/export/xlsx", headers=session)
    assert response.status_code == 400


def test_reconcile_without_files_is_rejected(client, session):
    response = client.post("/api/reconcile", headers=session, json={"mappings": []})
    assert response.status_code == 400


def test_upload_broker_only_cannot_reconcile(client, session, demo_files):
    upload = client.post(
        "/api/upload",
        headers=session,
        files={"files": ("broker.csv", demo_files["broker_register_sep2026.csv"], "text/csv")},
        data={"side": "broker"},
    )
    assert upload.status_code == 200
    response = client.post("/api/reconcile", headers=session, json={"mappings": []})
    assert response.status_code == 400
    assert "insurer statement" in response.json()["detail"]


def test_upload_rejects_unsupported_type(client, session):
    response = client.post(
        "/api/upload",
        headers=session,
        files={"files": ("summary.docx", b"random bytes", "application/msword")},
        data={"side": "statement"},
    )
    assert response.status_code == 200
    errors = response.json()["errors"]
    assert errors and "Unsupported" in errors[0]["error"]


def test_bad_mapping_is_caught_instead_of_producing_silent_nonsense(client, session):
    """Pointing Gross Premium at a text column must fail loudly, not reconcile to zero."""
    loaded = client.post("/api/demo/load", headers=session).json()
    broker = next(s for s in loaded["sources"] if s["side"] == "broker")
    currency_col = next(c["index"] for c in broker["columns"] if c["label"] == "Currency")
    response = client.post(
        "/api/reconcile",
        headers=session,
        json={"mappings": [{"id": broker["id"], "mapping": {"premium": currency_col}}]},
    )
    assert response.status_code == 400
    assert "No readable premiums" in response.json()["detail"]


def test_statement_with_no_policy_lines_is_rejected(client, session):
    """A header-only statement is a data error and is refused at upload."""
    csv_bytes = b"Policy No,Insured Name,Gross Premium,Commission Amount\n"
    response = client.post(
        "/api/upload",
        headers=session,
        files={"files": ("empty.csv", csv_bytes, "text/csv")},
        data={"side": "statement"},
    )
    assert response.status_code == 200
    errors = response.json()["errors"]
    assert errors and "no data rows" in errors[0]["error"]


def test_source_can_be_removed(client, session):
    loaded = client.post("/api/demo/load", headers=session).json()
    source_id = loaded["sources"][0]["id"]
    assert client.delete(f"/api/source/{source_id}", headers=session).status_code == 200
    assert client.post("/api/reset", headers=session).status_code == 200


def test_demo_file_download(client):
    response = client.get("/api/demo/file/GIG_statement_sep2026.xlsx")
    assert response.status_code == 200
    assert response.content[:2] == b"PK"
    assert client.get("/api/demo/file/nope.csv").status_code == 404
