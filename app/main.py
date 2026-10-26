"""ReconIQ web app - a local, zero-dependency-service prototype.

Everything stays in memory: files uploaded for a demo are parsed, reconciled and
discarded when the process stops. No database, no cloud, no external calls -
which is the only responsible way to run a demo with a broker's data.
"""

from __future__ import annotations

import csv
import io
import secrets
import time
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from fastapi import FastAPI, File, Form, Header, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel

from reconiq import __version__, demo
from reconiq.engine import Config, build_records, reconcile
from reconiq.ingest import IngestError, prepare
from reconiq.releases import RELEASES
from reconiq.quality import QUALITY
from reconiq.schema import describe_mapping, mapping_ready, normalize_mapping

BASE_DIR = Path(__file__).resolve().parent
STATIC_DIR = BASE_DIR / "static"
MAX_FILES = 8
SESSION_TTL_SECONDS = 12 * 60 * 60
MAX_SESSIONS = 40

app = FastAPI(title="ReconIQ", version=__version__, docs_url="/api/docs", redoc_url=None)


# --------------------------------------------------------------------------- #
# Session store - in-memory only, deliberately
# --------------------------------------------------------------------------- #
@dataclass
class Source:
    id: str
    filename: str
    side: str
    header: list[str]
    rows: list[list[str]]
    mapping: dict
    context: dict
    meta: dict
    skipped: list[dict] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


@dataclass
class Session:
    id: str
    created: float
    sources: dict[str, Source] = field(default_factory=dict)
    result: dict | None = None


SESSIONS: dict[str, Session] = {}


def _prune() -> None:
    now = time.time()
    for sid in [k for k, v in SESSIONS.items() if now - v.created > SESSION_TTL_SECONDS]:
        SESSIONS.pop(sid, None)
    while len(SESSIONS) > MAX_SESSIONS:
        oldest = min(SESSIONS.values(), key=lambda s: s.created)
        SESSIONS.pop(oldest.id, None)


def _session(session_id: str | None, create: bool = True) -> Session:
    _prune()
    if session_id and session_id in SESSIONS:
        return SESSIONS[session_id]
    if not create:
        raise HTTPException(status_code=404, detail="Session not found. Reload the page.")
    session = Session(id=secrets.token_urlsafe(12), created=time.time())
    SESSIONS[session.id] = session
    return session


# --------------------------------------------------------------------------- #
# Models
# --------------------------------------------------------------------------- #
class SourceRef(BaseModel):
    id: str
    side: str


class MappingOverride(BaseModel):
    id: str
    mapping: dict[str, int | None]


class ToleranceConfig(BaseModel):
    premium_abs_tol: float | None = None
    premium_pct_tol: float | None = None
    commission_abs_tol: float | None = None
    commission_pct_tol: float | None = None
    date_tol_days: int | None = None
    name_similarity_min: float | None = None


class ReconcileRequest(BaseModel):
    mappings: list[MappingOverride] = []
    config: ToleranceConfig = ToleranceConfig()


# --------------------------------------------------------------------------- #
# Routes
# --------------------------------------------------------------------------- #
@app.get("/api/health")
def health() -> dict:
    return {"status": "ok", "version": __version__, "sessions": len(SESSIONS)}


@app.get("/api/releases")
def release_history() -> dict:
    """Public product updates. Reading these never creates or changes a session."""
    return {"current_version": __version__, "releases": RELEASES}


@app.get("/api/quality")
def quality_overview() -> dict:
    """Public test catalogue and recorded verification snapshot; read-only."""
    return QUALITY


@app.post("/api/session")
def create_session() -> dict:
    session = _session(None)
    return {"session_id": session.id}


@app.get("/api/demo")
def demo_catalogue() -> dict:
    files = []
    for name, label, kind in demo.DEMO_FILES:
        files.append({"name": name, "label": label, "kind": kind})
    return {
        "broker": demo.BROKER_NAME,
        "period": demo.STATEMENT_PERIOD,
        "files": files,
        "note": (
            "Synthetic data built for this demo: one broker register and three insurer "
            "statements in three different formats, with a known set of exceptions planted."
        ),
    }


@app.get("/api/demo/file/{name}")
def demo_file(name: str) -> Response:
    allowed = {n: n for n, _, _ in demo.DEMO_FILES}
    if name not in allowed:
        raise HTTPException(status_code=404, detail="No such demo file.")
    payload = demo.build_all()[name]
    return Response(
        content=payload,
        media_type="application/octet-stream",
        headers={"Content-Disposition": f'attachment; filename="{name}"'},
    )


def _source_payload(source: Source, preview_rows: int = 8) -> dict:
    ready, missing = mapping_ready(source.mapping)
    return {
        "id": source.id,
        "filename": source.filename,
        "side": source.side,
        "header": source.header,
        "mapping": describe_mapping(source.mapping),
        "ready": ready,
        "missing": missing,
        "row_count": len(source.rows),
        "preview": source.rows[:preview_rows],
        "context": source.context,
        "meta": source.meta,
        "skipped": source.skipped[:35],
        "skipped_count": len(source.skipped),
        "notes": source.notes[:8],
        "columns": [
            {"index": idx, "label": label or f"(column {idx + 1})"}
            for idx, label in enumerate(source.header)
        ],
    }


def _store_source(session: Session, filename: str, data: bytes, side: str) -> Source:
    if side not in {"broker", "statement"}:
        raise HTTPException(status_code=400, detail="Unknown file role.")
    try:
        prepared = prepare(filename, data, side)
    except IngestError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc

    source = Source(
        id=secrets.token_hex(6),
        filename=filename,
        side=side,
        header=prepared["header"],
        rows=prepared["rows"],
        mapping=prepared["mapping"],
        context=prepared["context"],
        meta=prepared["meta"],
    )
    records, skipped, notes = build_records(source.id, side, source.rows, source.mapping)
    source.skipped = skipped
    source.notes = notes
    source.meta["record_count"] = len(records)
    session.sources[source.id] = source
    return source


@app.post("/api/upload")
async def upload(
    files: list[UploadFile] = File(...),
    side: str = Form(...),
    x_session: str | None = Header(default=None),
) -> dict:
    session = _session(x_session)
    if len(files) > MAX_FILES:
        raise HTTPException(status_code=400, detail=f"Upload at most {MAX_FILES} files at a time.")
    added, errors = [], []
    for upload_file in files:
        data = await upload_file.read()
        try:
            source = _store_source(session, upload_file.filename or "upload", data, side)
            added.append(_source_payload(source))
        except HTTPException as exc:
            errors.append({"filename": upload_file.filename, "error": exc.detail})
    return {"session_id": session.id, "sources": added, "errors": errors}


@app.post("/api/demo/load")
def load_demo(x_session: str | None = Header(default=None)) -> dict:
    session = _session(x_session)
    payloads = demo.build_all()
    added = []
    for name, _label, _kind in demo.DEMO_FILES:
        side = "broker" if name.startswith("broker") else "statement"
        source = _store_source(session, name, payloads[name], side)
        added.append(_source_payload(source))
    return {"session_id": session.id, "sources": added, "errors": []}


@app.delete("/api/source/{source_id}")
def delete_source(source_id: str, x_session: str | None = Header(default=None)) -> dict:
    session = _session(x_session, create=False)
    session.sources.pop(source_id, None)
    return {"ok": True}


@app.post("/api/reset")
def reset(x_session: str | None = Header(default=None)) -> dict:
    session = _session(x_session, create=False)
    session.sources.clear()
    session.result = None
    return {"ok": True}


def _apply_mapping(source: Source, override: dict[str, int | None]) -> None:
    mapping = normalize_mapping(source.mapping)
    for key, column in override.items():
        if key not in mapping:
            continue
        if column is None:
            mapping[key] = {"column": None, "confidence": 0, "header": None}
        else:
            column = int(column)
            if 0 <= column < len(source.header):
                mapping[key] = {
                    "column": column,
                    "confidence": 100,
                    "header": source.header[column],
                }
    source.mapping = mapping


def _config(payload: ToleranceConfig) -> Config:
    config = Config()

    def dec(value: float | None, fallback: Decimal) -> Decimal:
        if value is None:
            return fallback
        try:
            return Decimal(str(value)).quantize(Decimal("0.01"))
        except InvalidOperation:
            return fallback

    config.premium_abs_tol = dec(payload.premium_abs_tol, config.premium_abs_tol)
    config.premium_pct_tol = dec(payload.premium_pct_tol, config.premium_pct_tol)
    config.commission_abs_tol = dec(payload.commission_abs_tol, config.commission_abs_tol)
    config.commission_pct_tol = dec(payload.commission_pct_tol, config.commission_pct_tol)
    if payload.date_tol_days is not None:
        config.date_tol_days = max(0, min(int(payload.date_tol_days), 365))
    if payload.name_similarity_min is not None:
        config.name_similarity_min = max(0.5, min(float(payload.name_similarity_min), 1.0))
    return config


@app.post("/api/reconcile")
def run_reconciliation(
    request: ReconcileRequest, x_session: str | None = Header(default=None)
) -> dict:
    session = _session(x_session, create=False)
    if not session.sources:
        raise HTTPException(status_code=400, detail="Upload at least one file first.")

    overrides = {item.id: item.mapping for item in request.mappings}
    for source in session.sources.values():
        if source.id in overrides:
            _apply_mapping(source, overrides[source.id])

    broker_records = []
    broker_rows = 0
    statement_payloads = []
    unready: list[dict] = []

    for source in session.sources.values():
        ready, missing = mapping_ready(source.mapping)
        if not ready:
            unready.append({"filename": source.filename, "missing": missing})
            continue
        records, skipped, _notes = build_records(
            source.id, source.side, source.rows, source.mapping
        )
        source.skipped = skipped
        if records and not any(r.values.get("premium") is not None for r in records):
            raise HTTPException(
                status_code=400,
                detail=(
                    f"No readable premiums in {source.filename}. Check that Gross Premium is "
                    f"mapped to the column that actually holds the premium amount."
                ),
            )
        if source.side == "broker":
            broker_records.extend(records)
            broker_rows += len(records)
        else:
            statement_payloads.append(
                {
                    "source_id": source.id,
                    "label": source.filename,
                    "records": records,
                    "insurer_guess": source.context.get("insurer_guess"),
                }
            )

    if unready:
        raise HTTPException(
            status_code=400,
            detail=(
                "These files still need a Policy Number and Gross Premium column: "
                + "; ".join(f"{u['filename']} ({', '.join(u['missing'])})" for u in unready)
            ),
        )
    if not broker_records:
        raise HTTPException(
            status_code=400,
            detail="No broker policy file in this session - add your policy register first.",
        )
    if not statement_payloads:
        raise HTTPException(
            status_code=400,
            detail="Add at least one insurer statement to reconcile against the book.",
        )

    config = _config(request.config)
    started = time.perf_counter()
    try:
        result = reconcile(broker_records, statement_payloads, config)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    elapsed = round(time.perf_counter() - started, 3)

    labels = {source.id: source.filename for source in session.sources.values()}
    for finding in result["findings"]:
        finding["source_label"] = labels.get(finding["source_id"], "Broker book")
    for pair in result["matched_pairs"]:
        pair["source_label"] = labels.get(pair["source_id"], "")

    result["elapsed_seconds"] = elapsed
    result["generated_at"] = time.strftime("%Y-%m-%d %H:%M:%S")
    result["source_details"] = [
        {
            "id": source.id,
            "filename": source.filename,
            "side": source.side,
            "method": source.meta.get("method"),
            "confidence": source.meta.get("confidence"),
            "row_count": len(source.rows),
            "skipped_count": len(source.skipped),
            "warnings": source.meta.get("warnings") or [],
            "context": source.context,
        }
        for source in session.sources.values()
    ]
    result["data_quality"] = [
        {
            "filename": source.filename,
            "skipped": source.skipped[:20],
            "skipped_count": len(source.skipped),
            "notes": source.notes[:10],
        }
        for source in session.sources.values()
        if source.skipped or source.notes
    ]
    result["assumptions"] = [
        f"Tolerances: {result['kpis']['tolerance_summary']}.",
        "Dates are read day-first (DD/MM/YYYY), which is the UAE convention.",
        "Commission rates below 1 are treated as fractions (0.125 = 12.5%).",
        "Multi-currency lines are compared in AED at an indicative rate table "
        "(1 USD = 3.6725 AED). Real deployments would stamp the rate and date used.",
        f"Manual effort is estimated at {config.assume_seconds_per_line}s per line - "
        "an assumption to validate with the broker, not a measurement.",
        "Statements with no readable insurer are scoped by matching policy numbers first.",
        "This run used files held in memory only; nothing was written to disk or sent anywhere.",
    ]
    session.result = result
    return result


# --------------------------------------------------------------------------- #
# Export
# --------------------------------------------------------------------------- #
@app.get("/api/export/{fmt}")
def export(fmt: str, x_session: str | None = Header(default=None)) -> Response:
    session = _session(x_session, create=False)
    if not session.result:
        raise HTTPException(status_code=400, detail="Run a reconciliation before exporting.")
    stamp = time.strftime("%Y%m%d-%H%M")

    if fmt == "csv":
        buffer = io.StringIO()
        writer = csv.writer(buffer)
        writer.writerow(
            [
                "Severity", "Exception type", "Policy number", "Insured", "Insurer",
                "Field", "Broker value", "Statement value", "Variance (AED)",
                "Source file", "Detail", "Recommended action",
            ]
        )
        for finding in session.result["findings"]:
            writer.writerow(
                [
                    finding["severity"],
                    finding["category_label"],
                    finding["policy_no"],
                    finding["insured"],
                    finding["insurer"],
                    finding["field"],
                    finding["broker_value"],
                    finding["insurer_value"],
                    finding["variance_aed"] if finding["variance_aed"] is not None else "",
                    finding.get("source_label", ""),
                    finding["detail"],
                    finding["action"],
                ]
            )
        return Response(
            content=buffer.getvalue().encode("utf-8-sig"),
            media_type="text/csv",
            headers={
                "Content-Disposition": f'attachment; filename="reconiq-exceptions-{stamp}.csv"'
            },
        )

    if fmt != "xlsx":
        raise HTTPException(status_code=400, detail="Export format must be xlsx or csv.")

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    navy = "132A45"
    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor=navy)
    thin = Side(style="thin", color="D7DEE7")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    severity_fills = {
        "high": PatternFill("solid", fgColor="FBE3E3"),
        "medium": PatternFill("solid", fgColor="FCF0DA"),
        "low": PatternFill("solid", fgColor="E8F1FA"),
    }

    sheet = workbook.active
    sheet.title = "Exceptions"
    headers = [
        "Severity", "Exception type", "Policy number", "Insured", "Insurer", "Field",
        "Broker value", "Statement value", "Variance (AED)", "Source file",
        "Recommended action", "Detail",
    ]
    sheet.append(headers)
    for finding in session.result["findings"]:
        sheet.append(
            [
                finding["severity"].title(),
                finding["category_label"],
                finding["policy_no"],
                finding["insured"],
                finding["insurer"],
                finding["field"],
                finding["broker_value"],
                finding["insurer_value"],
                finding["variance_aed"],
                finding.get("source_label", ""),
                finding["action"],
                finding["detail"],
            ]
        )
    widths = [10, 26, 22, 32, 22, 18, 20, 20, 16, 28, 52, 72]
    for idx, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(idx)].width = width
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = border
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        severity = str(row[0].value).lower()
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=cell.column in (11, 12))
        if severity in severity_fills:
            row[0].fill = severity_fills[severity]
        if isinstance(row[8].value, (int, float)):
            row[8].number_format = '#,##0.00;[Red]-#,##0.00'
    sheet.row_dimensions[1].height = 28
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:L{max(sheet.max_row, 1)}"

    summary = workbook.create_sheet("Summary")
    kpis = session.result["kpis"]
    summary.append(["ReconIQ - reconciliation summary"])
    summary["A1"].font = Font(bold=True, size=14, color=navy)
    summary.append(["Generated", session.result.get("generated_at", "")])
    summary.append([])
    rows = [
        ("Broker policy lines", kpis["broker_records"]),
        ("Insurer statement lines", kpis["statement_records"]),
        ("Matched policies", kpis["matched"]),
        ("Clean matches (no exceptions)", kpis["clean_matches"]),
        ("Match rate %", kpis["match_rate"]),
        ("Exceptions raised", kpis["exceptions"]),
        ("High severity", kpis["by_severity"].get("high", 0)),
        ("Medium severity", kpis["by_severity"].get("medium", 0)),
        ("Low severity", kpis["by_severity"].get("low", 0)),
        ("Commission in broker book (AED)", kpis["broker_commission_aed"]),
        ("Commission on statements (AED)", kpis["statement_commission_aed"]),
        ("Net commission variance (AED)", kpis["commission_variance_aed"]),
        ("Unbilled commission - not on statement (AED)", kpis["unbilled_commission_aed"]),
        ("Unrecognised commission - not in book (AED)", kpis["unrecognised_commission_aed"]),
        ("Duplicate commission exposure (AED)", kpis["duplicate_exposure_aed"]),
        ("Commission on cancelled policies (AED)", kpis["cancelled_commission_aed"]),
        ("Total exception exposure (AED)", kpis["exposure_aed"]),
        ("Estimated manual review hours", kpis["manual_hours_estimate"]),
    ]
    for label, value in rows:
        summary.append([label, value])
    for row in summary.iter_rows(min_row=4, max_row=summary.max_row, max_col=2):
        row[0].font = Font(bold=True, size=10)
        if isinstance(row[1].value, float) and abs(row[1].value) > 1000:
            row[1].number_format = "#,##0.00"
    summary.column_dimensions["A"].width = 46
    summary.column_dimensions["B"].width = 22

    summary.append([])
    summary.append(["Assumptions and method"])
    summary.cell(row=summary.max_row, column=1).font = Font(bold=True, size=11, color=navy)
    for note in session.result.get("assumptions", []):
        summary.append([note])

    unmatched = workbook.create_sheet("Matched pairs")
    unmatched.append(
        ["Policy number", "Statement policy number", "Insured", "Insurer", "Premium",
         "Commission", "Match type", "Exceptions", "Statement file"]
    )
    for pair in session.result["matched_pairs"]:
        unmatched.append(
            [
                pair["policy_no"], pair["statement_policy_no"], pair["insured"],
                pair["insurer"], pair["premium"], pair["commission"], pair["status"],
                pair["finding_count"], pair.get("source_label", ""),
            ]
        )
    for cell in unmatched[1]:
        cell.font = header_font
        cell.fill = header_fill
    for idx, width in enumerate([24, 26, 32, 22, 16, 16, 18, 12, 30], start=1):
        unmatched.column_dimensions[get_column_letter(idx)].width = width
    unmatched.freeze_panes = "A2"

    buffer = io.BytesIO()
    workbook.save(buffer)
    return Response(
        content=buffer.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="reconiq-exceptions-{stamp}.xlsx"'},
    )


@app.get("/")
def index() -> FileResponse:
    return FileResponse(STATIC_DIR / "index.html")


app.mount("/static", StaticFiles(directory=STATIC_DIR), name="static")
