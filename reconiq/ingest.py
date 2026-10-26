"""File ingestion: CSV / TSV, Excel workbooks and PDF insurer statements.

PDF is the interesting one. Insurer statements arrive in two shapes:

1. ruled tables  -> pdfplumber's table extraction handles these well;
2. aligned text  -> columns drawn as fixed-width text with no lines at all.

We try ruled tables first, then a word-clustering fallback that rebuilds columns
from x-positions. Every result carries a method + confidence so the person
running the demo can *see* how the numbers were read before trusting them.
"""

from __future__ import annotations

import csv
import io
import re
import statistics
from dataclasses import dataclass, field

from .normalize import clean_text, parse_money, to_text
from .schema import FIELDS, detect_columns, header_row_index

MAX_ROWS = 20000
MAX_BYTES = 20 * 1024 * 1024

MONTHS = (
    "january|february|march|april|may|june|july|august|september|october|november|december"
    "|jan|feb|mar|apr|jun|jul|aug|sep|sept|oct|nov|dec"
)


@dataclass
class RawTable:
    rows: list[list[str]] = field(default_factory=list)
    meta: dict = field(default_factory=dict)


class IngestError(Exception):
    pass


def ingest(filename: str, data: bytes) -> RawTable:
    if len(data) > MAX_BYTES:
        raise IngestError(f"{filename} is larger than 20 MB - trim it before uploading.")
    lower = (filename or "").lower()
    if lower.endswith((".csv", ".txt", ".tsv")):
        table = _read_delimited(filename, data)
    elif lower.endswith((".xlsx", ".xlsm")):
        table = _read_xlsx(filename, data)
    elif lower.endswith(".pdf"):
        table = _read_pdf(filename, data)
    elif lower.endswith(".xls"):
        raise IngestError("Legacy .xls is not supported - save the file as .xlsx and re-upload.")
    else:
        raise IngestError(f"Unsupported file type for {filename}. Use CSV, XLSX or PDF.")

    if not table.rows:
        raise IngestError(f"No rows could be read from {filename}.")
    if len(table.rows) > MAX_ROWS:
        table.meta.setdefault("warnings", []).append(
            f"Only the first {MAX_ROWS:,} rows were read."
        )
        table.rows = table.rows[:MAX_ROWS]
    table.meta["filename"] = filename
    table.meta["bytes"] = len(data)
    return table


# --------------------------------------------------------------------------- #
# CSV / delimited text
# --------------------------------------------------------------------------- #
def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("latin-1", errors="replace")


def _read_delimited(filename: str, data: bytes) -> RawTable:
    text = _decode(data)
    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = "\t" if sample.count("\t") > sample.count(",") else ","
        if sample.count(";") > max(sample.count(","), sample.count("\t")):
            delimiter = ";"

    rows = [
        [clean_text(cell) for cell in row]
        for row in csv.reader(io.StringIO(text), delimiter=delimiter)
    ]
    width = max((len(r) for r in rows), default=0)
    rows = [r + [""] * (width - len(r)) for r in rows if any(clean_text(c) for c in r)]
    return RawTable(
        rows=rows,
        meta={
            "kind": "csv",
            "method": "delimited text",
            "delimiter": "TAB" if delimiter == "\t" else delimiter,
            "confidence": "high",
            "warnings": [],
        },
    )


# --------------------------------------------------------------------------- #
# Excel
# --------------------------------------------------------------------------- #
def _read_xlsx(filename: str, data: bytes) -> RawTable:
    from openpyxl import load_workbook

    warnings: list[str] = []
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - surfaced to the user
        raise IngestError(f"Could not open {filename}: {exc}") from exc

    best_sheet, best_rows = None, []
    for sheet in wb.worksheets:
        rows: list[list[str]] = []
        for raw_row in sheet.iter_rows(values_only=True):
            cells = [clean_text(c) for c in raw_row]
            if any(cells):
                rows.append(cells)
        width = max((len(r) for r in rows), default=0)
        rows = [r + [""] * (width - len(r)) for r in rows]
        if len(rows) > len(best_rows):
            best_sheet, best_rows = sheet.title, rows

    if len(wb.worksheets) > 1:
        others = [s.title for s in wb.worksheets if s.title != best_sheet]
        warnings.append(
            f"Workbook has {len(wb.worksheets)} sheets; read '{best_sheet}' "
            f"(most populated). Ignored: {', '.join(others[:5])}."
        )
    wb.close()
    return RawTable(
        rows=best_rows,
        meta={
            "kind": "xlsx",
            "sheet": best_sheet,
            "method": "excel workbook",
            "confidence": "high",
            "warnings": warnings,
        },
    )


# --------------------------------------------------------------------------- #
# PDF
# --------------------------------------------------------------------------- #
def _read_pdf(filename: str, data: bytes) -> RawTable:
    import pdfplumber

    warnings: list[str] = []
    ruled_rows: list[list[str]] = []
    text_rows: list[list[str]] = []
    page_count = 0

    try:
        pdf = pdfplumber.open(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise IngestError(
            f"Could not open the PDF ({exc}). If it is a scan, run OCR first or "
            "upload the CSV/Excel version of the statement."
        ) from exc

    page_text_parts: list[str] = []
    with pdf:
        page_count = len(pdf.pages)
        for page in pdf.pages[:60]:
            if len(page_text_parts) < 2:
                page_text_parts.append(page.extract_text() or "")
            for table in _page_tables(page):
                ruled_rows.extend(table)
            text_rows.extend(_page_text_rows(page))

    if _looks_like_statement(ruled_rows):
        rows, method, confidence = ruled_rows, "pdf ruled table", "high"
    elif _looks_like_statement(text_rows):
        rows, method, confidence = text_rows, "pdf text columns", "medium"
        warnings.append(
            "No table grid found in the PDF - columns were rebuilt from text "
            "positions. Check the preview before reconciling."
        )
    elif _looks_like_statement(ruled_rows):
        rows, method, confidence = ruled_rows, "pdf ruled table", "medium"
    else:
        pdf_text = ""
        with pdfplumber.open(io.BytesIO(data)) as probe:
            pdf_text = "\n".join((p.extract_text() or "") for p in probe.pages[:3])
        if len(pdf_text.strip()) < 40:
            raise IngestError(
                "This PDF has no extractable text - it is almost certainly a scan. "
                "OCR it first, or upload the CSV/Excel statement instead."
            )
        raise IngestError(
            "Could not find a policy table in this PDF. Try exporting the statement "
            "as CSV/Excel, or paste the rows into a spreadsheet."
        )

    rows = _dedupe_repeated_headers(rows)
    return RawTable(
        rows=rows,
        meta={
            "kind": "pdf",
            "pages": page_count,
            "method": method,
            "confidence": confidence,
            "warnings": warnings,
            "page_text": "\n".join(page_text_parts),
        },
    )


def _has_ruling(page) -> bool:
    """True if the page draws a real table grid.

    A single header underline is not a grid - treating it as one makes column
    detection worse than reading aligned text, so require several rules.
    """
    rules = len(page.lines) + len(page.rects) + len(page.curves)
    return rules >= 6


def _page_tables(page) -> list[list[list[str]]]:
    """Extract candidate tables from a page, strongest first.

    If the page draws no rules at all, pdfplumber's text strategy invents
    columns from whitespace and gets them subtly wrong - so we skip it and let
    the word-position rebuilder handle the page instead.
    """
    if not _has_ruling(page):
        return []
    strategies = (
        {"vertical_strategy": "lines", "horizontal_strategy": "lines"},
        {"vertical_strategy": "lines", "horizontal_strategy": "text"},
        {"vertical_strategy": "text", "horizontal_strategy": "text"},
    )
    best: list[list[str]] = []
    best_score = 0
    for settings in strategies:
        try:
            tables = page.extract_tables(settings) or []
        except Exception:  # noqa: BLE001 - pdfplumber can choke on odd geometry
            continue
        for table in tables:
            rows = [[clean_text(c) for c in row] for row in table if row]
            if not rows:
                continue
            width = max(len(r) for r in rows)
            rows = [r + [""] * (width - len(r)) for r in rows]
            score = _statement_score(rows)
            if score > best_score:
                best, best_score = rows, score
        if best_score >= 40:
            break
    return [best] if best else []


def _page_text_rows(page) -> list[list[str]]:
    """Rebuild columns from word x-positions for PDFs with no table lines.

    Words are first merged into cells (small intra-cell gaps), then cell starts
    are clustered into columns. Clustering raw words instead would split
    multi-word names and amounts across columns.
    """
    words = page.extract_words(use_text_flow=False, keep_blank_chars=False) or []
    if not words:
        return []

    lines: dict[float, list[dict]] = {}
    for word in words:
        bucket = None
        for existing in lines:
            if abs(existing - word["top"]) <= 3.2:
                bucket = existing
                break
        lines.setdefault(bucket if bucket is not None else round(word["top"], 1), []).append(word)
    ordered_lines = [sorted(lines[key], key=lambda w: w["x0"]) for key in sorted(lines)]

    cell_lines: list[list[tuple[str, float]]] = []
    for line in ordered_lines:
        cells: list[list[dict]] = []
        current: list[dict] = [line[0]]
        for previous, word in zip(line, line[1:]):
            if word["x0"] - previous["x1"] <= 7.5:
                current.append(word)
            else:
                cells.append(current)
                current = [word]
        cells.append(current)
        cell_lines.append(
            [(" ".join(w["text"] for w in cell), cell[0]["x0"]) for cell in cells]
        )

    anchors = _column_anchors([[start for _, start in cells] for cells in cell_lines], len(cell_lines))
    if not anchors:
        return []

    rows: list[list[str]] = []
    for cells in cell_lines:
        row = [""] * len(anchors)
        for text, start in cells:
            idx = min(range(len(anchors)), key=lambda i: abs(anchors[i] - start))
            row[idx] = clean_text(f"{row[idx]} {text}")
        if any(row):
            rows.append(row)
    return rows


def _column_anchors(starts_per_line: list[list[float]], line_count: int) -> list[float]:
    """Cluster cell left-edges across lines into stable column positions."""
    starts = sorted(value for line in starts_per_line for value in line)
    if not starts:
        return []
    clusters: list[list[float]] = [[starts[0]]]
    for value in starts[1:]:
        if value - clusters[-1][-1] <= 10:
            clusters[-1].append(value)
        else:
            clusters.append([value])
    min_hits = max(2, line_count // 5)
    anchors = [statistics.median(c) for c in clusters if len(c) >= min_hits]
    return sorted(anchors)


def _statement_score(rows: list[list[str]]) -> int:
    if not rows:
        return 0
    width = max(len(r) for r in rows)
    if width < 3 or len(rows) < 2:
        return 0
    numeric_cells = 0
    filled = 0
    for row in rows[:60]:
        for cell in row:
            if not clean_text(cell):
                continue
            filled += 1
            if parse_money(cell) is not None and not re.fullmatch(r"\d{1,2}", cell.strip()):
                numeric_cells += 1
    if filled == 0:
        return 0
    density = numeric_cells / filled
    return int(min(len(rows), 40) + width * 3 + density * 40)


def _looks_like_statement(rows: list[list[str]]) -> bool:
    return _statement_score(rows) >= 22


def _dedupe_repeated_headers(rows: list[list[str]]) -> list[list[str]]:
    """Drop header lines repeated at the top of continuation pages."""
    out: list[list[str]] = []
    seen_headers: list[list[str]] = []
    for row in rows:
        norm = [re.sub(r"[^a-z0-9]", "", clean_text(c).lower()) for c in row]
        joined = "".join(norm)
        if not joined:
            continue
        is_header = False
        for previous in seen_headers:
            if joined == previous:
                is_header = True
                break
        if is_header:
            continue
        alpha_cells = sum(1 for c in row if c and parse_money(c) is None)
        if alpha_cells >= max(3, len(row) // 2):
            seen_headers.append(joined)
            seen_headers = seen_headers[-4:]
        out.append(row)
    return out


# --------------------------------------------------------------------------- #
# Preparing a source for reconciliation
# --------------------------------------------------------------------------- #
def prepare(filename: str, data: bytes, side: str) -> dict:
    """Ingest a file and split it into context, header and data rows."""
    table = ingest(filename, data)
    rows = table.rows

    header_idx = header_row_index(rows)
    warnings = list(table.meta.get("warnings") or [])
    if header_idx is None:
        raise IngestError(
            f"Could not identify a header row in {filename}. The file needs column "
            "titles such as Policy No, Insured, Premium and Commission."
        )

    header = [clean_text(c) for c in rows[header_idx]]
    body = [r for r in rows[header_idx + 1 :] if any(clean_text(c) for c in r)]
    body = [r for r in body if not _is_total_row(r)]
    pre_rows = rows[:header_idx]
    context = _context_from(pre_rows, filename, table.meta.get("page_text") or "")

    if len(body) < 1:
        raise IngestError(f"{filename} has a header but no data rows underneath it.")

    header, body = _drop_empty_columns(header, body)
    mapping = detect_columns(header, body)
    return {
        "header": header,
        "rows": body,
        "pre_rows": pre_rows,
        "mapping": mapping,
        "context": context,
        "meta": {**table.meta, "warnings": warnings, "row_count": len(body), "side": side},
    }


def _drop_empty_columns(header: list[str], rows: list[list[str]]) -> tuple[list[str], list[list[str]]]:
    """Remove columns that carry no data at all - PDF column rebuilds invent these."""
    if not rows:
        return header, rows
    width = len(header)
    keep = [
        idx
        for idx in range(width)
        if clean_text(header[idx]) or any(idx < len(r) and clean_text(r[idx]) for r in rows)
    ]
    if len(keep) == width:
        return header, rows
    return (
        [header[i] for i in keep],
        [[row[i] if i < len(row) else "" for i in keep] for row in rows],
    )


def _is_total_row(row: list[str]) -> bool:
    first = clean_text(row[0]).lower() if row else ""
    joined = " ".join(clean_text(c).lower() for c in row[:3])
    return bool(re.match(r"^(total|totals|grand total|sub ?total|closing balance)", first)) or (
        "grand total" in joined
    )


def _context_from(pre_rows: list[list[str]], filename: str, page_text: str = "") -> dict:
    lines = [clean_text(c) for row in pre_rows for c in row if clean_text(c)]
    if not lines and page_text:
        lines = [clean_text(line) for line in page_text.splitlines() if clean_text(line)]
    text = " | ".join(lines)
    period = None
    match = re.search(rf"({MONTHS})\s*[-/,]?\s*(\d{{4}})", text, re.I)
    if match:
        period = f"{match.group(1).title()} {match.group(2)}"
    statement_no = None
    match = re.search(
        r"(?:statement|stmt)\s*(?:no\.?|number|#)?\s*[:#-]?\s*([A-Z0-9][A-Z0-9\-/]{4,})",
        text, re.I,
    )
    if match:
        statement_no = match.group(1)

    # Insurer is usually the first non-empty line of a statement.
    insurer_guess = None
    for line in lines[:4]:
        if len(line) > 3 and not re.search(r"\d{4}", line) and "statement" not in line.lower():
            insurer_guess = line
            break

    return {
        "title_lines": lines[:6],
        "period": period,
        "statement_no": statement_no,
        "insurer_guess": insurer_guess,
        "filename": filename,
    }
